"""Local extraction utilities for user documents and engineering drawings."""

import base64
import binascii
import io
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple
from xml.etree import ElementTree

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_FILES = 80
MAX_PDF_PAGES = 20

_ocr_engine = None


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR

        _ocr_engine = RapidOCR()
    return _ocr_engine


def _extract_image(image_bytes: bytes, page_number: int = 1) -> Tuple[str, List[Dict[str, Any]]]:
    from PIL import Image
    import numpy as np

    image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    result, _ = _get_ocr_engine()(np.asarray(image))
    items = []
    lines = []
    for row in result or []:
        box, text, confidence = row
        text = str(text).strip()
        if not text:
            continue
        points = [[round(float(point[0])), round(float(point[1]))] for point in box]
        items.append({
            "page": page_number,
            "text": text,
            "confidence": round(float(confidence), 3),
            "bbox": points,
        })
        lines.append(text)
    return "\n".join(lines), items


def _extract_docx(document_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(document_bytes)) as archive:
        xml_bytes = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml_bytes)
    return "\n".join(
        text for text in (element.text for element in root.iter() if element.tag.endswith("}t"))
        if text
    )


def _extract_xlsx(workbook_bytes: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    rows = []
    cell_count = 0
    for sheet in workbook.worksheets:
        rows.append(f"[SHEET] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
                cell_count += len(values)
                if cell_count >= 50000:
                    rows.append("[Extraction limit reached]")
                    return "\n".join(rows)
    return "\n".join(rows)


def _extract_pdf(pdf_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    import fitz

    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    lines = []
    ocr_items = []
    for page_index, page in enumerate(document[:MAX_PDF_PAGES], start=1):
        text = page.get_text("text").strip()
        if len(text) >= 24:
            lines.append(f"[PAGE {page_index}]\n{text}")
            continue
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        page_text, page_items = _extract_image(pixmap.tobytes("png"), page_index)
        if page_text:
            lines.append(f"[PAGE {page_index} OCR]\n{page_text}")
            ocr_items.extend(page_items)
    return "\n\n".join(lines), ocr_items


def _extract_member(name: str, file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(file_bytes)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}:
        return _extract_image(file_bytes)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(file_bytes), []
    if suffix == ".docx":
        return _extract_docx(file_bytes), []
    if suffix in {".txt", ".csv", ".md"}:
        return file_bytes.decode("utf-8", errors="replace"), []
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def extract_uploaded_document(file_name: str, data_url: str) -> Dict[str, Any]:
    """Decode a browser data URL, then pass its bytes through the shared extractor."""
    if "," not in data_url:
        raise ValueError("Upload must be a base64 data URL")
    try:
        file_bytes = base64.b64decode(data_url.split(",", 1)[1], validate=True)
    except (binascii.Error, ValueError) as error:
        raise ValueError("Uploaded file is not valid base64 data") from error
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("File exceeds the 50 MB limit")

    return extract_document_bytes(file_name, file_bytes)


def extract_document_bytes(file_name: str, file_bytes: bytes) -> Dict[str, Any]:
    """Extract document bytes. ZIP members are inspected in-memory, never written to disk."""
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("File exceeds the 50 MB limit")

    if Path(file_name).suffix.lower() != ".zip":
        text, items = _extract_member(file_name, file_bytes)
        return {
            "file_name": file_name,
            "text": text.strip(),
            "ocr_items": items,
            "extracted": bool(text.strip()),
            "source_count": 1,
        }

    collected_text = []
    collected_items = []
    count = 0
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
        for member in archive.infolist():
            if member.is_dir() or member.file_size > MAX_UPLOAD_BYTES:
                continue
            if count >= MAX_ARCHIVE_FILES:
                break
            try:
                text, items = _extract_member(member.filename, archive.read(member))
            except ValueError:
                continue
            count += 1
            if text.strip():
                collected_text.append(f"[SOURCE {member.filename}]\n{text.strip()}")
            collected_items.extend(items)

    combined = "\n\n".join(collected_text)
    return {
        "file_name": file_name,
        "text": combined,
        "ocr_items": collected_items,
        "extracted": bool(combined),
        "source_count": count,
    }